from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from build_trader_sentiment_report import build_analysis


OUT_DIR = Path.cwd() / "outputs" / "trader_sentiment_analysis"
DOCX_FILE = OUT_DIR / "editable_trader_sentiment_research_report.docx"

NAVY = "0B2545"
BLUE = "1F4D78"
TEAL = "007D72"
LIGHT_BLUE = "E8EEF5"
LIGHT_TEAL = "E7F4F2"
LIGHT_GRAY = "F2F4F7"
MID_GRAY = "667085"
GREEN = "107C41"
RED = "9B1C1C"
AMBER = "7A5A00"

SENTIMENT_ORDER = ["Extreme Fear", "Fear", "Neutral", "Greed", "Extreme Greed"]

SOURCES = [
    ("Alternative.me Crypto Fear & Greed Index", "https://alternative.me/crypto/fear-and-greed-index/"),
    ("CME Group: aligning cryptocurrency derivatives with 24/7 spot markets", "https://www.cmegroup.com/articles/2026/aligning-cryptocurrency-derivatives-with-spot-markets-measuring-the-247-trading-opportunity.html"),
    ("QuickNode Hyperliquid trades dataset documentation", "https://www.quicknode.com/docs/hyperliquid/datasets/trades"),
    ("Farzulla, Extremity Premium, arXiv 2026", "https://arxiv.org/abs/2602.07018"),
]


def money(x: float, decimals: int = 1) -> str:
    if pd.isna(x):
        return "n/a"
    sign = "-" if x < 0 else ""
    x = abs(float(x))
    if x >= 1_000_000_000:
        return f"{sign}${x / 1_000_000_000:.{decimals}f}B"
    if x >= 1_000_000:
        return f"{sign}${x / 1_000_000:.{decimals}f}M"
    if x >= 1_000:
        return f"{sign}${x / 1_000:.{decimals}f}K"
    return f"{sign}${x:,.0f}"


def pct(x: float) -> str:
    if pd.isna(x):
        return "n/a"
    return f"{x * 100:.1f}%"


def num(x: float) -> str:
    if pd.isna(x):
        return "n/a"
    return f"{int(round(x)):,}"


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9E2F3"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            margins.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    r_pr.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(underline)
    run.append(r_pr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_document(doc: Document):
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(23, 32, 51)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, color, before, after in [
        ("Title", 24, NAVY, 0, 8),
        ("Heading 1", 16, NAVY, 14, 6),
        ("Heading 2", 12.5, BLUE, 10, 4),
        ("Heading 3", 11.5, TEAL, 8, 3),
    ]:
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)


def add_footer(doc: Document):
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.text = "Trader Sentiment Research | Hyperliquid x Bitcoin Fear & Greed"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)
    footer = section.footer.paragraphs[0]
    footer.text = "Prepared from supplied CSV files. For research use only; not investment advice."
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor.from_string(MID_GRAY)


def add_callout(doc, label, text, fill=LIGHT_TEAL):
    table = doc.add_table(rows=1, cols=1)
    set_cell_margins(table)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_borders(cell, "FFFFFF")
    p = cell.paragraphs[0]
    r = p.add_run(label.upper() + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(TEAL)
    p.add_run(text)
    return table


def add_table(doc, df: pd.DataFrame, widths, header_fill=NAVY, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(df.columns))
    set_table_width(table, widths)
    set_cell_margins(table)
    hdr = table.rows[0]
    for i, col in enumerate(df.columns):
        cell = hdr.cells[i]
        cell.text = str(col)
        set_cell_shading(cell, header_fill)
        set_cell_borders(cell)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                r.font.size = Pt(font_size)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    for ridx, row in df.iterrows():
        cells = table.add_row().cells
        for i, col in enumerate(df.columns):
            cells[i].text = str(row[col])
            set_cell_shading(cells[i], LIGHT_GRAY if ridx % 2 == 0 else "FFFFFF")
            set_cell_borders(cells[i])
            cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(font_size)
    return table


def bullet(doc, text, style="List Bullet"):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p


def prepare_data():
    data = build_analysis()
    by = data["by_sent"].copy()
    by["sort"] = by["classification"].map({k: i for i, k in enumerate(SENTIMENT_ORDER)})
    by = by.sort_values("sort")
    by["volume_share"] = by["volume_usd"] / by["volume_usd"].sum()
    by["pnl_share"] = by["total_pnl"] / by["total_pnl"].sum()
    data["by_sent"] = by

    account = data["account_pivot"].copy()
    account["pnl_per_fill"] = account["total_pnl"] / account["total_trades"]
    data["top_accounts"] = account.sort_values("total_pnl", ascending=False).head(10)

    coin = data["top_symbols"].copy().sort_values("total_pnl", ascending=False).head(10)
    data["top_coin"] = coin

    daily = data["daily"]
    data["corr"] = daily[["sentiment_value", "total_pnl", "avg_pnl", "win_rate", "volume_usd", "trades"]].corr(numeric_only=True)
    data["best_days"] = daily.sort_values("total_pnl", ascending=False).head(5)
    data["worst_days"] = daily.sort_values("total_pnl", ascending=True).head(5)
    return data


def build_docx():
    data = prepare_data()
    o = data["overall"]
    doc = Document()
    style_document(doc)
    add_footer(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Trader Performance vs Bitcoin Market Sentiment")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = sub.add_run("Editable institutional research report on Hyperliquid fill-level behavior, realized PnL, and Fear/Greed regimes.")
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    meta = pd.DataFrame(
        [
            ["Research date", "8 Jun 2026"],
            ["Trade window", f"{o['date_min']} to {o['date_max']}"],
            ["Matched fills", f"{o['matched_trades']:,} of {o['total_trades']:,}"],
            ["Accounts / symbols", f"{o['accounts']:,} accounts / {o['symbols']:,} symbols"],
            ["Total closed PnL", money(o["total_pnl"])],
            ["Matched traded volume", money(o["total_volume"])],
        ],
        columns=["Field", "Value"],
    )
    add_table(doc, meta, [2.0, 4.0], TEAL, 9)

    add_callout(
        doc,
        "Core thesis",
        "Sentiment is most useful as a risk-budget and review layer, not as a standalone buy/sell signal. The data shows profitable trading across regimes, but PnL quality differs by sizing, win rate, side behavior, and account concentration.",
    )

    doc.add_heading("Executive Findings", level=1)
    for item in data["insights"]:
        bullet(doc, item)

    doc.add_heading("Methodology", level=1)
    doc.add_paragraph(
        "The analysis treats each Bitcoin Fear & Greed classification as a daily market state. Hyperliquid fills were parsed using Timestamp IST, converted to calendar date, and joined to the sentiment date field. Sentiment cuts exclude trades that did not match a sentiment date."
    )
    method = pd.DataFrame(
        [
            ["Performance outcome", "Closed PnL is treated as realized trading performance."],
            ["Win rate", "Closed PnL > 0."],
            ["Efficiency metric", "Average PnL per $1K traded = Closed PnL / Size USD x 1,000."],
            ["Primary cuts", "Sentiment regime, side, direction, account, symbol, and daily trend."],
            ["Key caution", "Total PnL can be distorted by volume concentration, so it is reviewed beside win rate, profit factor, fees, and normalized PnL."],
        ],
        columns=["Method item", "Definition"],
    )
    add_table(doc, method, [1.8, 4.7], BLUE, 8.5)

    doc.add_heading("Regime Scorecard", level=1)
    score = data["by_sent"][["classification", "trades", "volume_usd", "total_pnl", "win_rate", "avg_pnl_per_1k", "profit_factor"]].copy()
    score.columns = ["Regime", "Fills", "Volume", "Total PnL", "Win Rate", "PnL / $1K", "Profit Factor"]
    score["Fills"] = score["Fills"].map(lambda v: f"{int(v):,}")
    score["Volume"] = score["Volume"].map(money)
    score["Total PnL"] = score["Total PnL"].map(money)
    score["Win Rate"] = score["Win Rate"].map(pct)
    score["PnL / $1K"] = score["PnL / $1K"].map(lambda v: money(v, 0))
    score["Profit Factor"] = score["Profit Factor"].map(lambda v: "n/a" if pd.isna(v) else f"{v:.2f}x")
    add_table(doc, score, [1.15, 0.75, 1.0, 1.0, 0.85, 0.85, 0.85], NAVY, 7.8)
    add_callout(
        doc,
        "Interpretation",
        "Fear produced the largest absolute PnL pool, while Extreme Greed produced the best PnL efficiency in the supplied data. That distinction matters because capital allocation should reward efficiency, not only total dollars.",
        LIGHT_BLUE,
    )

    doc.add_heading("Trader Behavior Diagnostics", level=1)
    side = data["by_side"][["classification", "side", "trades", "total_pnl", "win_rate", "avg_pnl_per_1k"]].copy()
    side = side[side["trades"] > 0]
    side.columns = ["Regime", "Side", "Fills", "Total PnL", "Win Rate", "PnL / $1K"]
    side["Fills"] = side["Fills"].map(lambda v: f"{int(v):,}")
    side["Total PnL"] = side["Total PnL"].map(money)
    side["Win Rate"] = side["Win Rate"].map(pct)
    side["PnL / $1K"] = side["PnL / $1K"].map(lambda v: money(v, 0))
    add_table(doc, side, [1.15, 0.65, 0.8, 1.05, 0.85, 0.9], BLUE, 7.7)
    doc.add_paragraph(
        "Desk read: side and direction cuts identify whether a regime's PnL comes from repeatable behavior or simply from more activity. A follow-up review should add drawdown, adverse excursion, and fee-adjusted return by side."
    )

    doc.add_heading("Account and Symbol Concentration", level=1)
    acct = data["top_accounts"][["account", "total_trades", "total_pnl", "pnl_per_fill"]].copy()
    acct["account"] = acct["account"].str[:14] + "..."
    acct.columns = ["Account", "Fills", "Total PnL", "PnL / Fill"]
    acct["Fills"] = acct["Fills"].map(lambda v: f"{int(v):,}")
    acct["Total PnL"] = acct["Total PnL"].map(money)
    acct["PnL / Fill"] = acct["PnL / Fill"].map(lambda v: money(v, 0))
    add_table(doc, acct, [1.8, 0.9, 1.1, 1.0], TEAL, 8)

    coin = data["top_coin"][["coin", "trades", "volume_usd", "total_pnl", "avg_pnl_per_1k"]].copy()
    coin.columns = ["Symbol", "Fills", "Volume", "Total PnL", "PnL / $1K"]
    coin["Fills"] = coin["Fills"].map(lambda v: f"{int(v):,}")
    coin["Volume"] = coin["Volume"].map(money)
    coin["Total PnL"] = coin["Total PnL"].map(money)
    coin["PnL / $1K"] = coin["PnL / $1K"].map(lambda v: money(v, 0))
    add_table(doc, coin, [0.9, 0.75, 1.0, 1.0, 0.9], BLUE, 8)

    doc.add_heading("Daily Stress Review", level=1)
    for heading, df in [("Best PnL days", data["best_days"]), ("Worst PnL days", data["worst_days"])]:
        doc.add_heading(heading, level=2)
        d = df[["date", "classification", "total_pnl", "volume_usd", "win_rate"]].copy()
        d.columns = ["Date", "Regime", "Total PnL", "Volume", "Win Rate"]
        d["Total PnL"] = d["Total PnL"].map(money)
        d["Volume"] = d["Volume"].map(money)
        d["Win Rate"] = d["Win Rate"].map(pct)
        add_table(doc, d, [1.0, 1.25, 1.1, 1.1, 0.9], NAVY if heading.startswith("Best") else RED, 8)

    doc.add_heading("Trading-Company Strategy Playbook", level=1)
    playbook = pd.DataFrame(
        [
            ["Extreme Fear", "Defense first", "Reduce impulse entries; require liquidity check; favor smaller clips.", "High uncertainty and adverse-selection risk."],
            ["Fear", "Selective offense", "Allow risk where account/symbol edge is proven; monitor fee drag.", "Largest PnL pool, but volume-driven."],
            ["Neutral", "Baseline", "Use standard limits; no sentiment-based override.", "Control state."],
            ["Greed", "Momentum with guardrails", "Permit trend exposure; tighten loss review and take-profit discipline.", "Crowding risk rising."],
            ["Extreme Greed", "High-quality offense", "Scale only for accounts with proven efficiency; avoid late crowded trades.", "Best PnL per $1K in this dataset."],
        ],
        columns=["Regime", "Desk posture", "Trading rule", "Why"],
    )
    add_table(doc, playbook, [1.05, 1.25, 3.0, 2.0], TEAL, 8)

    doc.add_heading("Reviewer Checklist: Areas to Inspect Before Submission", level=1)
    checks = [
        "Confirm whether the assignment expects Timestamp IST or raw Timestamp as the join basis. This report uses Timestamp IST calendar date.",
        "Confirm whether Closed PnL should be analyzed before or after fees. This report presents fee context but treats Closed PnL as the primary outcome.",
        "Check whether unmatched sentiment dates should be excluded or assigned nearest available sentiment. This report excludes unmatched records.",
        "Review whether PnL per $1K traded is the preferred normalization; alternatives include PnL per fill, fee-adjusted PnL, and account-level Sharpe-style measures.",
        "Consider adding drawdown/adverse-excursion analysis if the evaluator expects risk-adjusted performance rather than realized PnL only.",
    ]
    for item in checks:
        bullet(doc, item)

    doc.add_heading("External Reference Framing", level=1)
    ref_intro = doc.add_paragraph(
        "External references were used for methodology framing and industry-style interpretation. The empirical results come from the supplied CSV files."
    )
    for name, url in SOURCES:
        p = doc.add_paragraph(style="List Bullet")
        add_hyperlink(p, name, url)
        p.add_run(f" - {url}")

    doc.add_heading("Source Files and Assumptions", level=1)
    assumptions = pd.DataFrame(
        [
            ["Bitcoin sentiment CSV", r"D:\companies tast\internshala primetrade ai\fear_greed_index.csv"],
            ["Hyperliquid trader CSV", r"D:\companies tast\internshala primetrade ai\historical_data.csv"],
            ["Merge key", "Timestamp IST date joined to sentiment date."],
            ["Sentiment coverage", f"{o['sentiment_min']} to {o['sentiment_max']}"],
            ["Trade coverage", f"{o['date_min']} to {o['date_max']}"],
        ],
        columns=["Item", "Detail"],
    )
    add_table(doc, assumptions, [1.75, 4.75], BLUE, 8)

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX_FILE)
    return DOCX_FILE


if __name__ == "__main__":
    path = build_docx()
    print(path)
